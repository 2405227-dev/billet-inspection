import os
import csv
import openpyxl
from pathlib import Path
from datetime import datetime
import pytest

from backend.services.exports.pdf_exporter import generate_individual_billet_pdf, generate_batch_audit_pdf
from backend.services.exports.excel_exporter import export_audit_excel
from backend.services.exports.csv_exporter import (
    export_inspections_csv,
    export_defects_csv,
    export_eca_csv,
    export_paut_csv,
    export_measurements_csv,
    export_audit_csv,
)

@pytest.fixture
def mock_passport_inspection_record():
    return {
        "id": "INS-2026-TEST001",
        "billet_id": "BILLET-00023",
        "heat_number": "H7821",
        "batch_id": "BATCH-24A",
        "steel_grade": "IS 2830 / Fe 500D",
        "timestamp": datetime(2026, 10, 4, 8, 21, 3),
        "status": "REJECT",
        "decision_reason": "Surface crack indication at 742 mm correlated with ECA; PAUT volumetric echo at 4400 mm (depth 75 mm).",
        "why_failed_explanation": "Surface crack at 742 mm · ECA impedance anomaly · PAUT core porosity at 75 mm depth.",
        "length_mm": 6002.0,
        "width_mm": 150.2,
        "height_mm": 150.1,
        "diameter_mm": None,
        "dimension_status": "PASS",
        "dimension_deviation_pct": 0.03,
        "encoder_entry_position_mm": 0.0,
        "encoder_exit_position_mm": 6002.0,
        "conveyor_speed_mm_s": 850.0,
        "coverage_top": "INSPECTED",
        "coverage_bottom": "NOT_INSPECTED",
        "coverage_left": "INSPECTED",
        "coverage_right": "INSPECTED",
        "coverage_front": "INSPECTED",
        "coverage_back": "INSPECTED",
        "defect_count": 1,
        "ocr_text": "HEAT-H7821 | BATCH-24A",
        "ocr_confidence": 0.985,
        "barcode_value": "BILLET-00023-H7821",
        "barcode_type": "DATA_MATRIX",
        "production_line_id": "LINE-BILLET-01",
        "plant_id": "PLANT-RAIPUR-01",
        "operator_id": "OP-9812",
        "model_version": "surface-defect-ensemble-seg:v3.2.0",
        "calibration_version": "CALIB-CAM01-2026-v4",
        "processing_latency_ms": 42.5,
        "audit_hash": "a1b2c3d4e5f67890123456789abcdef0123456789abcdef0123456789abcdef0",
        "previous_audit_hash": "0000000000000000000000000000000000000000000000000000000000000000",
        "evidence_image": "/storage/inspection-images/INS-2026-TEST001_evidence.jpg",
        "dimensions": [
            {
                "dimension_name": "length",
                "measured_value_mm": 6002.0,
                "nominal_mm": 6000.0,
                "expected_min_mm": 5800.0,
                "expected_max_mm": 6200.0,
                "deviation_pct": 0.03,
                "status": "PASS",
                "camera_id": "CAM-01",
            },
            {
                "dimension_name": "width",
                "measured_value_mm": 150.2,
                "nominal_mm": 150.0,
                "expected_min_mm": 147.0,
                "expected_max_mm": 153.0,
                "deviation_pct": 0.13,
                "status": "PASS",
                "camera_id": "CAM-01",
            },
            {
                "dimension_name": "height",
                "measured_value_mm": 150.1,
                "nominal_mm": 150.0,
                "expected_min_mm": 147.0,
                "expected_max_mm": 153.0,
                "deviation_pct": 0.07,
                "status": "PASS",
                "camera_id": "CAM-01",
            },
        ],
        "defects": [
            {
                "id": "DEF-0007",
                "defect_type": "surface_crack",
                "severity": "MAJOR",
                "confidence": 0.94,
                "surface_face": "TOP",
                "position_from_front_mm": 742.0,
                "position_from_back_mm": 5260.0,
                "encoder_pos_mm": 742.0,
                "x_pos_px": 280,
                "y_pos_px": 410,
                "length_mm": 38.0,
                "width_mm": 0.7,
                "area_mm2": 26.6,
                "orientation_deg": 12.5,
                "crop_path": "/storage/defect-crops/DEF-0007.jpg",
            }
        ],
        "ocr_results": [
            {
                "raw_text": "HEAT-H7821 | BATCH-24A",
                "parsed_heat_number": "H7821",
                "parsed_batch_id": "BATCH-24A",
                "confidence": 0.985,
                "bbox": [320, 410, 260, 55],
                "engine_name": "PaddleOCR",
            }
        ],
        "eca_indications": [
            {
                "channel_idx": 7,
                "signal_amplitude_mv": 296.0,
                "phase_deg": 88.5,
                "encoder_position_mm": 650.0,
                "position_from_front_mm": 650.0,
                "indication_type": "SURFACE_CRACK_IMPEDANCE",
                "severity": "MAJOR",
                "confidence": 0.92,
                "is_simulated": True,
            }
        ],
        "paut_indications": [
            {
                "channel_idx": 32,
                "depth_mm": 75.0,
                "amplitude_pct": 68.0,
                "amplitude": 68.0,
                "encoder_position_mm": 4400.0,
                "position_from_front_mm": 4400.0,
                "indication_type": "INTERNAL_CORE_POROSITY",
                "severity": "CRITICAL",
                "confidence": 0.94,
                "is_simulated": True,
            }
        ],
    }

def test_passport_pdf_generation(tmp_path, mock_passport_inspection_record):
    pdf_path = tmp_path / "passport_BILLET-00023.pdf"
    res = generate_individual_billet_pdf(mock_passport_inspection_record, pdf_path)
    assert res.exists()
    assert res.stat().st_size > 1000 # Verified PDF output generated

def test_11_sheet_excel_export(tmp_path, mock_passport_inspection_record):
    excel_path = tmp_path / "billet_inspection_audit.xlsx"
    res = export_audit_excel([mock_passport_inspection_record], excel_path)
    assert res.exists()

    wb = openpyxl.load_workbook(str(res))
    expected_sheets = [
        "01_Inspection_Summary",
        "02_Billet_Dimensions",
        "03_Surface_Defects",
        "04_ECA_Indications",
        "05_PAUT_Indications",
        "06_OCR_Results",
        "07_3D_Measurements",
        "08_Sensor_Correlation",
        "09_Inspection_Coverage",
        "10_Decision",
        "11_Audit_Trail",
    ]
    for s_name in expected_sheets:
        assert s_name in wb.sheetnames

    # Verify Billet ID is present across all sheets as primary linking key
    for s_name in expected_sheets:
        ws = wb[s_name]
        assert ws.cell(row=1, column=1).value == "Billet ID"
        assert ws.cell(row=2, column=1).value == "BILLET-00023"

def test_csv_exports_consistency(tmp_path, mock_passport_inspection_record):
    billets_csv = tmp_path / "billets.csv"
    defects_csv = tmp_path / "defects.csv"
    eca_csv = tmp_path / "eca_indications.csv"
    paut_csv = tmp_path / "paut_indications.csv"
    measurements_csv = tmp_path / "measurements.csv"
    audit_csv = tmp_path / "audit.csv"

    export_inspections_csv([mock_passport_inspection_record], billets_csv)
    export_defects_csv([mock_passport_inspection_record], defects_csv)
    export_eca_csv([mock_passport_inspection_record], eca_csv)
    export_paut_csv([mock_passport_inspection_record], paut_csv)
    export_measurements_csv([mock_passport_inspection_record], measurements_csv)
    export_audit_csv([mock_passport_inspection_record], audit_csv)

    # 1. Check billets.csv
    with open(billets_csv, mode="r", encoding="utf-8") as f:
        reader = list(csv.DictReader(f))
        assert len(reader) == 1
        assert reader[0]["billet_id"] == "BILLET-00023"
        assert reader[0]["heat_number"] == "H7821"
        assert float(reader[0]["length_mm"]) == 6002.0
        assert reader[0]["status"] == "REJECT"

    # 2. Check defects.csv
    with open(defects_csv, mode="r", encoding="utf-8") as f:
        reader = list(csv.DictReader(f))
        assert len(reader) == 1
        assert reader[0]["billet_id"] == "BILLET-00023"
        assert reader[0]["defect_id"] == "DEF-0007"
        assert float(reader[0]["position_from_front_mm"]) == 742.0
        assert float(reader[0]["length_mm"]) == 38.0
        assert float(reader[0]["width_mm"]) == 0.7

    # 3. Check eca_indications.csv
    with open(eca_csv, mode="r", encoding="utf-8") as f:
        reader = list(csv.DictReader(f))
        assert len(reader) == 1
        assert reader[0]["billet_id"] == "BILLET-00023"
        assert float(reader[0]["signal_amplitude_mv"]) == 296.0
        assert float(reader[0]["encoder_position_mm"]) == 650.0

    # 4. Check paut_indications.csv
    with open(paut_csv, mode="r", encoding="utf-8") as f:
        reader = list(csv.DictReader(f))
        assert len(reader) == 1
        assert reader[0]["billet_id"] == "BILLET-00023"
        assert float(reader[0]["depth_mm"]) == 75.0
        assert float(reader[0]["amplitude_pct"]) == 68.0
        assert float(reader[0]["encoder_position_mm"]) == 4400.0

    # 5. Check measurements.csv
    with open(measurements_csv, mode="r", encoding="utf-8") as f:
        reader = list(csv.DictReader(f))
        assert len(reader) == 3 # length, width, height
        assert reader[0]["billet_id"] == "BILLET-00023"
        assert reader[0]["dimension_name"] == "Length"
        assert float(reader[0]["measured_value_mm"]) == 6002.0
